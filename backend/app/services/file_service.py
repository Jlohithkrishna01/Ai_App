import os
import uuid
from pathlib import Path
from typing import Tuple, Optional
from fastapi import UploadFile, HTTPException
from PIL import Image
from pypdf import PdfReader
import docx

from app.config import settings

ALLOWED_EXTENSIONS = {
    ".pdf", ".docx", ".doc", ".txt", ".csv", ".md", ".json",
    ".png", ".jpg", ".jpeg", ".webp"
}

ALLOWED_AVATAR_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}

MAX_FILE_BYTES = settings.MAX_FILE_SIZE_MB * 1024 * 1024

def validate_uploaded_file(file: UploadFile) -> Tuple[str, str]:
    if not file.filename:
        raise HTTPException(status_code=400, detail="Uploaded file must have a filename.")
    
    ext = Path(file.filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported file type '{ext}'. Allowed types: PDF, DOCX, TXT, CSV, MD, JSON, PNG, JPG, WEBP."
        )
    return file.filename, ext

def extract_text_from_file(file_path: Path, ext: str) -> str:
    extracted_text = ""
    try:
        if ext == ".pdf":
            reader = PdfReader(str(file_path))
            pages_text = []
            for i, page in enumerate(reader.pages):
                text = page.extract_text()
                if text:
                    pages_text.append(f"--- Page {i + 1} ---\n{text.strip()}")
            extracted_text = "\n\n".join(pages_text)

        elif ext in [".docx", ".doc"]:
            doc = docx.Document(str(file_path))
            full_text = []
            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        full_text.append(row_text)
            extracted_text = "\n".join(full_text)

        elif ext in [".txt", ".csv", ".md", ".json"]:
            try:
                with open(file_path, "r", encoding="utf-8") as f:
                    extracted_text = f.read()
            except UnicodeDecodeError:
                with open(file_path, "r", encoding="latin-1") as f:
                    extracted_text = f.read()

        elif ext in [".png", ".jpg", ".jpeg", ".webp"]:
            # For image files, extract dimensions and basic metadata
            with Image.open(file_path) as img:
                extracted_text = f"[Uploaded Image: {file_path.name}, Format: {img.format}, Dimensions: {img.width}x{img.height}]"

    except Exception as e:
        print(f"[FileService] Error extracting text from {file_path}: {e}")
        extracted_text = f"[Could not fully extract text from document: {str(e)}]"

    return extracted_text.strip()

async def save_uploaded_file(file: UploadFile) -> dict:
    filename, ext = validate_uploaded_file(file)
    
    contents = await file.read()
    if len(contents) > MAX_FILE_BYTES:
        raise HTTPException(
            status_code=400,
            detail=f"File exceeds maximum allowed size of {settings.MAX_FILE_SIZE_MB}MB."
        )
    
    unique_name = f"{uuid.uuid4().hex}_{Path(filename).name}"
    save_path = settings.FILES_DIR / unique_name
    
    with open(save_path, "wb") as f:
        f.write(contents)
    
    extracted_text = extract_text_from_file(save_path, ext)
    file_url = f"/uploads/files/{unique_name}"
    
    return {
        "file_name": filename,
        "saved_name": unique_name,
        "file_url": file_url,
        "file_type": ext.lstrip("."),
        "file_size": len(contents),
        "extracted_text": extracted_text
    }

async def save_avatar_file(file: UploadFile) -> str:
    if not file.filename:
        raise HTTPException(status_code=400, detail="Avatar file must have a filename.")
    ext = Path(file.filename).suffix.lower()
    if ext not in ALLOWED_AVATAR_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail=f"Invalid avatar format. Allowed: PNG, JPG, JPEG, WEBP."
        )
    
    contents = await file.read()
    if len(contents) > 5 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="Avatar file size exceeds 5MB limit.")
    
    unique_name = f"avatar_{uuid.uuid4().hex}.jpg"
    save_path = settings.AVATARS_DIR / unique_name
    
    # Process with PIL to standardize to 256x256 square avatar
    from io import BytesIO
    try:
        with Image.open(BytesIO(contents)) as img:
            img = img.convert("RGB")
            # Crop to square
            min_dim = min(img.width, img.height)
            left = (img.width - min_dim) // 2
            top = (img.height - min_dim) // 2
            right = left + min_dim
            bottom = top + min_dim
            img_cropped = img.crop((left, top, right, bottom))
            img_resized = img_cropped.resize((256, 256), Image.Resampling.LANCZOS)
            img_resized.save(save_path, "JPEG", quality=90)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to process image: {str(e)}")

    return f"/uploads/avatars/{unique_name}"
